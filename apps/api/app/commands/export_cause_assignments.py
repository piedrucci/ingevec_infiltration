"""Export group/category/cause assignments from the María Platias workbook sheets to a new private JSON file."""
from __future__ import annotations

import argparse
import json
import os
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ASSIGNMENT_SHEET = "MARIA PLATIAS"
CATALOG_SHEET = "Catalogo_causas"
INCIDENT_SHEET = "IncidenciasCausas"


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _number(value: Any) -> str | None:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def _split(value: Any) -> list[str]:
    text = _text(value)
    return [part.strip() for part in text.split("/") if part.strip()] if text else []


def _records(sheet) -> list[dict[str, Any]]:
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None) or ()
    index = {name: position for position, name in enumerate(header) if name is not None}
    return [
        {name: row[position] if position < len(row) else None for name, position in index.items()}
        for row in rows
    ]


def build_assignments(workbook) -> tuple[list[dict[str, Any]], list[str]]:
    """Return JSON rows and blocking issues; any issue means the export must not be written.

    A single row category receives every cause. With several categories, each cause takes the one row
    category allowed by the cause catalog, falling back to the per-cause incident sheet.
    """
    catalog: dict[str, set[str]] = defaultdict(set)
    for row in workbook[CATALOG_SHEET].iter_rows(min_row=2, values_only=True):
        if row and _text(row[0]) and len(row) > 1 and _text(row[1]):
            catalog[_text(row[0])].add(_text(row[1]))
    incidents: dict[tuple[str | None, str], str | None] = {}
    for record in _records(workbook[INCIDENT_SHEET]):
        cause = _text(record.get("Analisis causa de Fallas"))
        if cause:
            incidents[(_number(record.get("ID")), cause)] = _text(record.get("Item Asociado"))

    rows: list[dict[str, Any]] = []
    issues: list[str] = []
    for record in _records(workbook[ASSIGNMENT_SHEET]):
        notes = _text(record.get("Obs"))
        if not notes:
            continue
        row_id = _number(record.get("ID"))
        group = _text(record.get("Categoría"))
        categories = _split(record.get("Item"))
        causes = _split(record.get("Analisis causa de Fallas"))
        if causes and not categories:
            issues.append(f"ID {row_id}: causes without category: {notes}")
        associated = []
        for cause in dict.fromkeys(causes):
            if len(categories) == 1:
                category = categories[0]
            else:
                candidates = catalog.get(cause, set()) & set(categories)
                fallback = incidents.get((row_id, cause))
                category = next(iter(candidates)) if len(candidates) == 1 else fallback if fallback in categories else None
            if category is None:
                if categories:
                    issues.append(f"ID {row_id}: cannot resolve category of {cause!r} among {categories}: {notes}")
                continue
            incident_category = incidents.get((row_id, cause))
            if incident_category is not None and incident_category != category:
                issues.append(f"ID {row_id}: {cause!r} resolved to {category!r} but {INCIDENT_SHEET} says {incident_category!r}")
            associated.append({"group": group, "category": category, "cause": cause})
        rows.append({"project_id": _number(record.get("N° Obra")), "notes": notes, "associated_causes": associated})

    duplicates = sorted(note for note, count in Counter(row["notes"] for row in rows).items() if count > 1)
    issues.extend(f"duplicate notes: {note}" for note in duplicates)
    return rows, issues


def summarize(rows: list[dict[str, Any]]) -> dict[str, int]:
    associations = [cause for row in rows for cause in row["associated_causes"]]
    return {
        "rows": len(rows),
        "associations": len(associations),
        "rows_without_causes": sum(1 for row in rows if not row["associated_causes"]),
        "associations_without_group": sum(1 for cause in associations if cause["group"] is None),
    }


def write_new(path: Path, rows: list[dict[str, Any]]) -> None:
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(fd, "w", encoding="utf-8") as output:
        json.dump({"rows": rows}, output, ensure_ascii=False, indent=2)
        output.write("\n")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True, help="New private JSON file (never overwritten)")
    args = parser.parse_args()
    if args.output.exists():
        raise SystemExit(f"Refusing to overwrite {args.output}")
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    rows, issues = build_assignments(workbook)
    if issues:
        print(json.dumps({"status": "blocked", "issues": issues}, ensure_ascii=False, indent=2))
        raise SystemExit(1)
    write_new(args.output, rows)
    print(json.dumps({"status": "written", "output": str(args.output), **summarize(rows)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
