import hashlib
import json
import re
import unicodedata
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import (
    Classification, DivisionManager, ExcelImport, ExcelSourceRow, ItemType,
    Location, PostventaItem, Project, ProjectAdmin, ProjectManager,
    Subcontractor, Supervisor, Typology,
)

SHEET = "Año 2026"
REQUIRED_HEADERS = {
    "N° Obra", "Proyecto", "Tipologia Proyecto", "Ubicación Proyecto",
    "Supervisor Pvta", "Clasificación", "Item", "Obs", "Fecha Solicitud",
}


def norm(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return re.sub(r"[^a-z0-9]+", "", text.encode("ascii", "ignore").decode().lower())


def clean(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for parser in (
            lambda candidate: datetime.strptime(candidate, "%d-%m-%Y").date(),
            lambda candidate: datetime.fromisoformat(candidate).date(),
            lambda candidate: date.fromisoformat(candidate),
        ):
            try:
                return parser(text)
            except ValueError:
                continue
        return None
    return None


def _raw_value(value: object) -> object:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _value(raw_cells: dict, header: str, occurrence: int = 1) -> object:
    matches = [cell.get("value") for cell in raw_cells["cells"] if norm(cell.get("header")) == norm(header)]
    return matches[occurrence - 1] if len(matches) >= occurrence else None


def stage_workbook(db: Session, content: bytes, filename: str) -> tuple[ExcelImport, bool]:
    digest = hashlib.sha256(content).hexdigest()
    existing = db.scalar(select(ExcelImport).where(ExcelImport.file_hash == digest))
    if existing is not None:
        return existing, False

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    if SHEET not in workbook.sheetnames:
        raise ValueError(f"Sheet {SHEET!r} not found")
    worksheet = workbook[SHEET]
    rows = worksheet.iter_rows(values_only=False)
    header_cells = next(rows, None)
    if header_cells is None:
        raise ValueError("The source sheet is empty")
    headers = [cell.value for cell in header_cells]
    present = {str(header).strip() for header in headers if header is not None}
    missing = sorted(REQUIRED_HEADERS - present)
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    imported = ExcelImport(
        original_filename=filename, file_hash=digest, status="STAGING",
        source_sheet=SHEET, row_count=0,
    )
    db.add(imported)
    db.flush()
    staged = 0
    for row_number, row in enumerate(rows, start=2):
        if not any(cell.value is not None for cell in row):
            continue
        raw_cells = {
            "cells": [
                {"position": index + 1, "header": _raw_value(headers[index]), "value": _raw_value(cell.value)}
                for index, cell in enumerate(row)
            ]
        }
        encoded = json.dumps(raw_cells, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
        db.add(ExcelSourceRow(
            excel_import_id=imported.id, sheet_name=SHEET, row_number=row_number,
            raw_cells=raw_cells, row_hash=hashlib.sha256(encoded).hexdigest(),
            normalization_status="PENDING",
        ))
        staged += 1
    imported.row_count = staged
    imported.status = "STAGED"
    db.commit()
    return imported, True


def _catalog(db: Session, model, value: object, caches: dict):
    name = clean(value)
    if name is None:
        return None
    cache = caches.setdefault(model, {norm(item.name): item for item in db.scalars(select(model)).all()})
    if norm(name) in cache:
        return cache[norm(name)]
    item = model(name=name)
    db.add(item)
    db.flush()
    cache[norm(name)] = item
    return item


def _normalize_row(db: Session, source: ExcelSourceRow, caches: dict) -> None:
    existing_item = db.scalar(select(PostventaItem).where(PostventaItem.source_row_id == source.id))
    if existing_item is not None:
        source.normalization_status = "NORMALIZED"
        source.normalization_error = None
        return
    raw = source.raw_cells
    typology = _catalog(db, Typology, _value(raw, "Tipologia Proyecto"), caches)
    location = _catalog(db, Location, _value(raw, "Ubicación Proyecto"), caches)
    supervisor = _catalog(db, Supervisor, _value(raw, "Supervisor Pvta"), caches)
    classification = _catalog(db, Classification, _value(raw, "Clasificación"), caches)
    # The workbook contains two Item columns; the second is the business item type.
    item_type = _catalog(db, ItemType, _value(raw, "Item", occurrence=2), caches)
    subcontractor = _catalog(db, Subcontractor, _value(raw, "Subcontrato"), caches)
    work_number = clean(_value(raw, "N° Obra"))
    project_name = clean(_value(raw, "Proyecto"))
    notes = clean(_value(raw, "Obs"))
    required = {
        "N° Obra": work_number, "Proyecto": project_name, "Tipologia Proyecto": typology,
        "Ubicación Proyecto": location, "Supervisor Pvta": supervisor,
        "Clasificación": classification, "Item": item_type, "Obs": notes,
    }
    missing = [name for name, value in required.items() if value is None]
    if missing:
        raise ValueError(f"Missing required values: {', '.join(missing)}")

    division_name = _value(raw, "GD")
    manager_name = _value(raw, "GP")
    admin_name = _value(raw, "AD")
    project_admin = None
    if division_name and manager_name:
        division = _catalog(db, DivisionManager, division_name, caches)
        manager = next((m for m in db.scalars(select(ProjectManager).where(ProjectManager.division_manager_id == division.id)).all() if norm(m.name) == norm(manager_name)), None)
        if manager is None:
            manager = ProjectManager(name=clean(manager_name), division_manager_id=division.id)
            db.add(manager)
            db.flush()
        if admin_name:
            project_admin = next((a for a in db.scalars(select(ProjectAdmin).where(ProjectAdmin.project_manager_id == manager.id)).all() if norm(a.name) == norm(admin_name)), None)
            if project_admin is None:
                project_admin = ProjectAdmin(name=clean(admin_name), project_manager_id=manager.id)
                db.add(project_admin)
                db.flush()

    project = db.get(Project, work_number)
    if project is None:
        project = Project(
            id=work_number, name=project_name, typology_id=typology.id,
            location_id=location.id, supervisor_id=supervisor.id,
            municipal_reception_date=as_date(_value(raw, "Recep. Municipal")),
            project_admin_id=project_admin.id if project_admin else None,
        )
        db.add(project)
        db.flush()
    elif project.project_admin_id is None and project_admin is not None:
        project.project_admin_id = project_admin.id
    db.add(PostventaItem(
        source_row_id=source.id, project_id=project.id, classification_id=classification.id,
        item_type_id=item_type.id, notes=notes,
        request_date=as_date(_value(raw, "Fecha Solicitud")),
        subcontractor_id=subcontractor.id if subcontractor else None,
        handled_by=clean(_value(raw, "Gestionado por")),
    ))
    source.normalization_status = "NORMALIZED"
    source.normalization_error = None


def normalize_import(db: Session, import_id, batch_size: int = 50, max_rows: int | None = None) -> dict:
    imported = db.get(ExcelImport, import_id)
    if imported is None:
        raise ValueError("Excel import not found")
    imported.status = "NORMALIZING"
    db.commit()
    processed = 0
    caches: dict = {}
    while max_rows is None or processed < max_rows:
        limit = min(batch_size, max_rows - processed) if max_rows is not None else batch_size
        batch = db.scalars(
            select(ExcelSourceRow)
            .where(ExcelSourceRow.excel_import_id == import_id, ExcelSourceRow.normalization_status == "PENDING")
            .order_by(ExcelSourceRow.row_number).limit(limit)
        ).all()
        if not batch:
            break
        for source in batch:
            try:
                _normalize_row(db, source, caches)
            except ValueError as exc:
                source.normalization_status = "ERROR"
                source.normalization_error = str(exc)
        db.commit()
        processed += len(batch)
    pending = db.scalar(select(func.count()).select_from(ExcelSourceRow).where(ExcelSourceRow.excel_import_id == import_id, ExcelSourceRow.normalization_status == "PENDING")) or 0
    errors = db.scalar(select(func.count()).select_from(ExcelSourceRow).where(ExcelSourceRow.excel_import_id == import_id, ExcelSourceRow.normalization_status == "ERROR")) or 0
    normalized = db.scalar(select(func.count()).select_from(ExcelSourceRow).where(ExcelSourceRow.excel_import_id == import_id, ExcelSourceRow.normalization_status == "NORMALIZED")) or 0
    imported.status = "COMPLETED_WITH_ERRORS" if pending == 0 and errors else "COMPLETED" if pending == 0 else "NORMALIZING"
    db.commit()
    return {"processed": processed, "pending": pending, "normalized": normalized, "errors": errors, "status": imported.status}


def backfill_dates(db: Session, import_id) -> dict:
    rows = db.scalars(
        select(ExcelSourceRow)
        .where(ExcelSourceRow.excel_import_id == import_id, ExcelSourceRow.normalization_status == "NORMALIZED")
        .order_by(ExcelSourceRow.row_number)
    ).all()
    items_by_source = {
        item.source_row_id: item
        for item in db.scalars(select(PostventaItem).where(PostventaItem.source_row_id.is_not(None))).all()
    }
    project_dates: dict[str, date] = {}
    updated_items = 0
    for source in rows:
        item = items_by_source.get(source.id)
        if item is None:
            continue
        request_date = as_date(_value(source.raw_cells, "Fecha Solicitud"))
        if item.request_date != request_date:
            item.request_date = request_date
            updated_items += 1
        reception_date = as_date(_value(source.raw_cells, "Recep. Municipal"))
        if reception_date is not None and item.project_id not in project_dates:
            project_dates[item.project_id] = reception_date
    updated_projects = 0
    for project_id, reception_date in project_dates.items():
        project = db.get(Project, project_id)
        if project is not None and project.municipal_reception_date != reception_date:
            project.municipal_reception_date = reception_date
            updated_projects += 1
    db.commit()
    return {"updated_projects": updated_projects, "updated_items": updated_items}


def backfill_project_admins(db: Session, import_id) -> dict:
    rows = db.scalars(
        select(ExcelSourceRow)
        .where(ExcelSourceRow.excel_import_id == import_id, ExcelSourceRow.normalization_status == "NORMALIZED")
        .order_by(ExcelSourceRow.row_number)
    ).all()
    caches: dict = {}
    updated = 0
    for source in rows:
        raw = source.raw_cells
        work_number = clean(_value(raw, "N° Obra"))
        division_name = clean(_value(raw, "GD"))
        manager_name = clean(_value(raw, "GP"))
        admin_name = clean(_value(raw, "AD"))
        if not all((work_number, division_name, manager_name, admin_name)):
            continue
        project = db.get(Project, work_number)
        if project is None or project.project_admin_id is not None:
            continue
        division = _catalog(db, DivisionManager, division_name, caches)
        manager = next((m for m in db.scalars(select(ProjectManager).where(ProjectManager.division_manager_id == division.id)).all() if norm(m.name) == norm(manager_name)), None)
        if manager is None:
            manager = ProjectManager(name=manager_name, division_manager_id=division.id)
            db.add(manager)
            db.flush()
        admin = next((a for a in db.scalars(select(ProjectAdmin).where(ProjectAdmin.project_manager_id == manager.id)).all() if norm(a.name) == norm(admin_name)), None)
        if admin is None:
            admin = ProjectAdmin(name=admin_name, project_manager_id=manager.id)
            db.add(admin)
            db.flush()
        project.project_admin_id = admin.id
        updated += 1
    db.commit()
    return {"updated_projects": updated}


def import_workbook(db: Session, content: bytes, filename: str, row_limit: int | None = None) -> tuple[ExcelImport, bool]:
    imported, created = stage_workbook(db, content, filename)
    if created or imported.status in {"STAGED", "NORMALIZING"}:
        normalize_import(db, imported.id, max_rows=row_limit)
        db.refresh(imported)
    return imported, created


def storage_key(content: bytes, name: str) -> str:
    safe_name = re.sub(r"[^A-Za-z0-9._-]+", "-", name).strip("-") or "source.xlsx"
    return f"imports/excel/{hashlib.sha256(content).hexdigest()}/{safe_name}"
