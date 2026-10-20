import json

import pytest
from openpyxl import Workbook

from app.commands.export_cause_assignments import build_assignments, summarize, write_new

HEADER = ["ID", "Analisis causa de Fallas", "Categoría", "N° Obra", "Item", "Obs"]


def workbook(rows, incidents=(), catalog=()):
    book = Workbook()
    sheet = book.active
    sheet.title = "MARIA PLATIAS"
    sheet.append(HEADER)
    for row in rows:
        sheet.append(row)
    incident_sheet = book.create_sheet("IncidenciasCausas")
    incident_sheet.append(["ID", "Analisis causa de Fallas", "Item Asociado"])
    for row in incidents:
        incident_sheet.append(row)
    catalog_sheet = book.create_sheet("Catalogo_causas")
    catalog_sheet.append(["Causa", "Item Asociado"])
    for row in catalog:
        catalog_sheet.append(row)
    return book


def test_single_category_receives_every_cause_and_numbers_become_strings():
    rows, issues = build_assignments(workbook([[1.0, "Sellos cortados / PVC inferior", "Ejecución", 712.0, "Ventana", " Depto A19 "]]))
    assert issues == []
    assert rows == [{"project_id": "712", "notes": "Depto A19", "associated_causes": [
        {"group": "Ejecución", "category": "Ventana", "cause": "Sellos cortados"},
        {"group": "Ejecución", "category": "Ventana", "cause": "PVC inferior"},
    ]}]


def test_several_categories_use_catalog_then_incident_sheet():
    book = workbook(
        [[2.0, "PVC inferior / Mala ejecución EIFS / Sellos cortados", "Ejecución", 712.0, "Ventana / Fachada", "Antepecho"]],
        incidents=[[2.0, "Sellos cortados", "Ventana"]],
        catalog=[["PVC inferior", "Ventana"], ["Mala ejecución EIFS", "Fachada"], ["Sellos cortados", "Ventana"], ["Sellos cortados", "Hojalatería"]],
    )
    rows, issues = build_assignments(book)
    assert issues == []
    assert [(c["category"], c["cause"]) for c in rows[0]["associated_causes"]] == [
        ("Ventana", "PVC inferior"), ("Fachada", "Mala ejecución EIFS"), ("Ventana", "Sellos cortados")]


def test_rows_without_group_or_causes_are_kept_empty():
    rows, issues = build_assignments(workbook([[3.0, None, None, 700.0, "Ventana", "Sin análisis"], [4.0, "X", "Diseño", 1.0, "Ventana", None]]))
    assert issues == []
    assert rows == [{"project_id": "700", "notes": "Sin análisis", "associated_causes": []}]
    assert summarize(rows)["rows_without_causes"] == 1


@pytest.mark.parametrize("rows, incidents, catalog, message", [
    ([[5.0, "Desconocida", "Ejecución", 1.0, "Ventana / Fachada", "A"]], [], [], "cannot resolve category"),
    ([[6.0, "PVC inferior", "Ejecución", 1.0, None, "B"]], [], [], "causes without category"),
    ([[7.0, "PVC inferior", "Ejecución", 1.0, "Ventana", "C"]], [[7.0, "PVC inferior", "Fachada"]], [], "says 'Fachada'"),
    ([[8.0, None, None, 1.0, None, "D"], [9.0, None, None, 2.0, None, "D"]], [], [], "duplicate notes"),
])
def test_contradictions_block_the_export(rows, incidents, catalog, message):
    _, issues = build_assignments(workbook(rows, incidents, catalog))
    assert any(message in issue for issue in issues)


def test_output_is_new_private_file(tmp_path):
    path = tmp_path / "out.json"
    write_new(path, [{"project_id": "1", "notes": "A", "associated_causes": []}])
    assert json.loads(path.read_text(encoding="utf-8"))["rows"][0]["notes"] == "A"
    assert path.stat().st_mode & 0o777 == 0o600
    with pytest.raises(FileExistsError):
        write_new(path, [])
